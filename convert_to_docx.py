"""
Convert ANALYSIS_DOCUMENTATION.md to a formatted MS Word document.
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Color palette ────────────────────────────────────────────────
DARK_BLUE     = RGBColor(0x1F, 0x49, 0x7D)   # heading H1
MID_BLUE      = RGBColor(0x2E, 0x74, 0xB5)   # heading H2
ACCENT_BLUE   = RGBColor(0x2F, 0x96, 0xD5)   # heading H3
DARK_GRAY     = RGBColor(0x26, 0x26, 0x26)   # body text
TABLE_HEADER  = RGBColor(0x1F, 0x49, 0x7D)   # table header bg
CODE_BG       = RGBColor(0xF3, 0xF3, 0xF3)   # code block bg
CODE_FG       = RGBColor(0x17, 0x17, 0x17)   # code text
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_bg(cell, color_hex: str):
    """Set table cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_table_borders(table):
    """Add thin borders to all table cells."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "BFBFBF")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def add_horizontal_rule(doc):
    """Add a styled horizontal separator line."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)
    return p


def apply_inline_formats(para, text: str):
    """
    Parse inline markdown: **bold**, `code`, and plain text —
    and add formatted runs to the paragraph.
    """
    # Split on **bold** and `code`
    pattern = re.compile(r'(\*\*.*?\*\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            if part:
                para.add_run(part)


def add_code_block(doc, code_text: str):
    """Add a code block paragraph with monospaced font and gray background."""
    lines = code_text.strip("\n").split("\n")
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        # Light gray background via paragraph shading
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F3F3F3")
        pPr.append(shd)
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        run.font.color.rgb = CODE_FG
    # spacer after code block
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)


def parse_markdown_table(lines, start_idx):
    """
    Parse a markdown table starting at lines[start_idx].
    Returns (headers, rows, end_idx).
    """
    headers = [h.strip() for h in lines[start_idx].strip("|").split("|")]
    # skip separator line (start_idx+1)
    rows = []
    i = start_idx + 2
    while i < len(lines):
        line = lines[i].strip()
        if not line.startswith("|"):
            break
        row = [c.strip() for c in line.strip("|").split("|")]
        rows.append(row)
        i += 1
    return headers, rows, i - 1


def build_word_table(doc, headers, rows):
    """Build a styled Word table from parsed markdown table data."""
    col_count = max(len(headers), max((len(r) for r in rows), default=1))
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_borders(table)

    # Header row
    hdr_row = table.rows[0]
    for ci, htext in enumerate(headers):
        if ci >= col_count:
            break
        cell = hdr_row.cells[ci]
        set_cell_bg(cell, "1F497D")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(htext)
        run.bold = True
        run.font.color.rgb = WHITE
        run.font.size = Pt(9)
        run.font.name = "Calibri"

    # Data rows
    for ri, row in enumerate(rows):
        tr = table.rows[ri + 1]
        bg = "FFFFFF" if ri % 2 == 0 else "EEF3FA"
        for ci in range(col_count):
            cell = tr.cells[ci]
            set_cell_bg(cell, bg)
            p = cell.paragraphs[0]
            cell_text = row[ci] if ci < len(row) else ""
            apply_inline_formats(p, cell_text)
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Calibri"
    doc.add_paragraph()  # spacer


def configure_styles(doc):
    """Set up document-level styles and page layout."""
    section = doc.sections[0]
    section.page_width  = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK_GRAY
    normal.paragraph_format.space_after = Pt(6)

    # Heading 1
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = DARK_BLUE
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after  = Pt(6)

    # Heading 2
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(15)
    h2.font.bold = True
    h2.font.color.rgb = MID_BLUE
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after  = Pt(4)

    # Heading 3
    h3 = doc.styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = ACCENT_BLUE
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after  = Pt(3)

    # Heading 4
    h4 = doc.styles["Heading 4"]
    h4.font.name = "Calibri"
    h4.font.size = Pt(11)
    h4.font.bold = True
    h4.font.color.rgb = DARK_GRAY
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after  = Pt(2)

    # List Bullet
    try:
        lb = doc.styles["List Bullet"]
        lb.font.name = "Calibri"
        lb.font.size = Pt(10.5)
        lb.paragraph_format.left_indent   = Inches(0.25)
        lb.paragraph_format.space_after   = Pt(3)
    except Exception:
        pass


def add_cover_page(doc):
    """Add a styled cover page."""
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("F&O Analyzer")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(36)
    title_run.font.bold = True
    title_run.font.color.rgb = DARK_BLUE

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Analysis Strategy Documentation")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(20)
    sub_run.font.color.rgb = MID_BLUE

    doc.add_paragraph()
    add_horizontal_rule(doc)

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_run = desc_p.add_run(
        "NIFTY & SENSEX Futures & Options Analysis Dashboard\n"
        "FastAPI + Zerodha KiteConnect\n\n"
        "Covers 19+ Analysis Modules:\n"
        "Black-Scholes Greeks · Implied Volatility · PCR · Max Pain\n"
        "GEX · IV Surface · MTF Signals · Strategy Payoff · Backtesting"
    )
    desc_run.font.name = "Calibri"
    desc_run.font.size = Pt(11)
    desc_run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    desc_p.paragraph_format.space_before = Pt(10)

    doc.add_paragraph()
    add_horizontal_rule(doc)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run("Lot Sizes: NIFTY = 65  |  SENSEX = 20\nVersion 1.0.0  |  June 2026")
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    doc.add_page_break()


def convert_md_to_docx(md_path: str, docx_path: str):
    doc = Document()
    configure_styles(doc)
    add_cover_page(doc)

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]

        # ── Skip the very first H1 (used as cover title already) ──
        if line.startswith("# ") and i < 5:
            i += 1
            continue

        # ── Horizontal rule ─────────────────────────────────────────
        if line.strip() in ("---", "***", "___"):
            add_horizontal_rule(doc)
            i += 1
            continue

        # ── Headings ────────────────────────────────────────────────
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 4")
            apply_inline_formats(p, line[5:].strip())
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            apply_inline_formats(p, line[4:].strip())
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            apply_inline_formats(p, line[3:].strip())
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Heading 1")
            apply_inline_formats(p, line[2:].strip())
            i += 1
            continue

        # ── Code block (``` ... ```) ─────────────────────────────────
        if line.strip().startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, "\n".join(code_lines))
            i += 1
            continue

        # ── Markdown table ───────────────────────────────────────────
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"\|[\s\-|]+\|", lines[i + 1]):
            headers, rows, end_idx = parse_markdown_table(lines, i)
            build_word_table(doc, headers, rows)
            i = end_idx + 1
            continue

        # ── Bullet list ──────────────────────────────────────────────
        if re.match(r"^(\s*[-*] |\s*\d+\. )", line):
            text = re.sub(r"^\s*[-*] ", "", line)
            text = re.sub(r"^\s*\d+\. ", "", text)
            indent_level = len(line) - len(line.lstrip())
            style = "List Bullet"
            p = doc.add_paragraph(style=style)
            if indent_level > 2:
                p.paragraph_format.left_indent = Inches(0.5)
            apply_inline_formats(p, text.strip())
            i += 1
            continue

        # ── Blank line ───────────────────────────────────────────────
        if line.strip() == "":
            # Don't add excessive blank paragraphs
            i += 1
            continue

        # ── Regular paragraph ────────────────────────────────────────
        p = doc.add_paragraph(style="Normal")
        apply_inline_formats(p, line.strip())
        i += 1

    # ── Footer on all pages ──────────────────────────────────────────
    for section in doc.sections:
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_para.add_run("F&O Analyzer — Analysis Documentation  |  NIFTY & SENSEX Dashboard  |  v1.0.0")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(docx_path)
    print(f"[OK] Saved: {docx_path}")


if __name__ == "__main__":
    import os
    base = r"D:\Sensex-Nifty-FutureOptions"
    convert_md_to_docx(
        md_path   = os.path.join(base, "ANALYSIS_DOCUMENTATION.md"),
        docx_path = os.path.join(base, "FO_Analyzer_Analysis_Documentation.docx"),
    )
